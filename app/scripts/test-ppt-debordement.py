"""Calibration de verifier_debordements_texte (pptx_deck.py).

Les sous-titres de colonne (9 pt, boite 3,75 x 0,22 in) sont rendus sur UNE
ligne par PowerPoint (mesure 2026-10-15) : ils ne doivent plus lever de faux
positif, mais un texte reellement trop long dans la meme boite doit toujours
etre signale. Usage : python test-ppt-debordement.py
"""
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import pptx_deck as D
from pptx import Presentation

echecs = 0


def check(cond, msg):
    global echecs
    if not cond:
        echecs += 1
    print(("  ok   " if cond else "  FAIL ") + msg)


SOUS_TITRES = [
    "Forte dispersion des réponses — sujets à clarifier",
    "Maturité la plus basse — leviers de progrès prioritaires",
    "Maturité la plus haute — points d'appui à valoriser",
    "Dispersion la plus faible des réponses — consensus fort",
]


def constats(texte, w=3.75, h=0.22, size=9):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    D.add_text(s, 0.5, 0.5, w, h, [(texte, {"size": size})])
    return D.verifier_debordements_texte(prs)


print("Sous-titres nominaux : aucun faux positif")
for t in SOUS_TITRES:
    check(constats(t) == [], f"pas de constat — {t[:35]}…")

print("Vrai debordement : toujours signale")
vrai = ("Un sous-titre volontairement beaucoup trop long qui se replie sur "
        "plusieurs lignes et sortirait réellement de sa boîte de colonne, "
        "bien au-delà d'une seule ligne de texte à neuf points.")
check(len(constats(vrai)) == 1, "texte de ~190 car. dans la boite de sous-titre flagué")
check(len(constats("Texte moyen " * 12, w=2.0, h=0.2, size=12)) == 1,
      "boite etroite et basse flaguee")

print("Deck nominal complet : aucun constat")
src = open(os.path.join(HERE, "test-export-ppt.py"), encoding="utf8").read().split("def main():")[0]
ns = {"__file__": os.path.join(HERE, "test-export-ppt.py")}
exec(compile(src, ns["__file__"], "exec"), ns)
gen = ns["gen"]
dep = ns["bloc_equipe"]("DSI Paiements", [2.0, 1.8, 1.6, 1.9], avec_comp=False)
dep["type"] = "departement"
dep["nbEquipes"] = 2
data = {"couverture": {"titre": "R", "sousTitre": "D", "date": "24/06/2026"},
        "blocs": [dep, ns["bloc_equipe"]("Squad Paiement", [2.7, 2.0, 1.4, 2.3])]}
tmp = tempfile.mkdtemp(prefix="test-debord-")
prs, _ = gen.construire(data, gen.TEMPLATE, os.path.join(tmp, "d.pptx"))
d = D.verifier_debordements_texte(prs)
check(d == [], f"deck nominal sans avertissement — {d}")

print(f"\n{echecs} echec(s)")
sys.exit(1 if echecs else 0)
